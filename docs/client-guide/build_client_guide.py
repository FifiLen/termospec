from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "Instrukcja-Nuxt-Studio-produkty-Termospec.docx"
SCREENSHOTS = ROOT / "docs" / "client-guide" / "screenshots"
LOGO = ROOT / "public" / "termospec-logo.png"

BRAND = "9A0B2E"
INK = "120B38"
MUTED = "6D6675"
LINE = "E6DDE2"
SURFACE = "F7F4F1"
WHITE = "FFFFFF"
BLUE = "2E74B5"


def set_run_font(run, name="Calibri", size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def configure_table_row(row, *, repeat_header=False):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    if repeat_header:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=9, color=MUTED)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BRAND)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)
    c_nv_pr = inline_shape._inline.graphic.graphicData.pic.nvPicPr.cNvPr
    c_nv_pr.set("name", title)
    c_nv_pr.set("title", title)
    c_nv_pr.set("descr", description)


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BRAND, 18, 10),
        "Heading 2": (13, BRAND, 14, 7),
        "Heading 3": (12, INK, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(10)


def configure_running_header_footer(doc: Document):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("TERMOSPEC  /  INSTRUKCJA NUxT STUDIO")
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("Termospec  •  20.08.2026  •  "), size=8.5, color=MUTED)
    add_page_number(p)


def set_keep_with_next(paragraph, enabled=True):
    paragraph.paragraph_format.keep_with_next = enabled


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None, italic=False):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]), italic=italic)
    else:
        set_run_font(p.add_run(text), italic=italic)
    return p


def add_bullets(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_run_font(p.add_run(item))


def create_numbering_instance(doc: Document, style_id="ListNumber") -> int:
    """Create a fresh Word numbering instance so each procedure starts at 1."""
    numbering = doc.part.numbering_part.element
    abstract_id = None

    for abstract in numbering.findall(qn("w:abstractNum")):
        for node in abstract.iter(qn("w:pStyle")):
            if node.get(qn("w:val")) == style_id:
                abstract_id = abstract.get(qn("w:abstractNumId"))
                break
        if abstract_id is not None:
            break

    if abstract_id is None:
        raise RuntimeError(f"No numbering definition found for style {style_id}")

    existing_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    num_id = max(existing_ids, default=0) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), abstract_id)
    num.append(abstract_num_id)

    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)

    return num_id


def apply_numbering(paragraph, num_id: int):
    paragraph_properties = paragraph._p.get_or_add_pPr()
    existing = paragraph_properties.find(qn("w:numPr"))
    if existing is not None:
        paragraph_properties.remove(existing)

    num_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_properties.extend([level, number])
    paragraph_properties.append(num_properties)


def add_numbered(doc: Document, items: Iterable[str]):
    num_id = create_numbering_instance(doc)
    for item in items:
        p = doc.add_paragraph(style="List Number")
        apply_numbering(p, num_id)
        set_run_font(p.add_run(item))


def add_callout(doc: Document, label: str, text: str, *, fill=SURFACE, accent=BRAND):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.2
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    set_run_font(p.add_run(f"{label}: "), bold=True, color=accent)
    set_run_font(p.add_run(text))
    return p


def add_screenshot(doc: Document, filename: str, caption: str, alt: str, width=6.5):
    path = SCREENSHOTS / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(0)
    set_keep_with_next(p)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    add_alt_text(shape, caption, alt)
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    return cap


def add_field_table(doc: Document, rows: list[tuple[str, str]], header=("Pole", "Jak je wypełnić")):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2700, 6660])
    hdr = table.rows[0].cells
    configure_table_row(table.rows[0], repeat_header=True)
    for idx, text in enumerate(header):
        set_cell_shading(hdr[idx], "E8EEF5")
        p = hdr[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(p.add_run(text), size=10, bold=True, color=INK)
    for label, detail in rows:
        cells = table.add_row().cells
        configure_table_row(table.rows[-1])
        for cell in cells:
            set_cell_margins(cell)
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        p0.paragraph_format.line_spacing = 1.1
        set_run_font(p0.add_run(label), size=10, bold=True, color=BRAND)
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.1
        set_run_font(p1.add_run(detail), size=10)
    return table


def add_page_break(doc: Document):
    doc.add_page_break()


def add_safe_page_break(doc: Document, extra_space=0):
    """Start a new page without LibreOffice collapsing the next block into the top margin."""
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(extra_space)
    p.paragraph_format.line_spacing = Pt(1)
    set_run_font(p.add_run("\u200b"), size=1, color=WHITE)


def build():
    doc = Document()
    style_document(doc)
    configure_running_header_footer(doc)
    doc.core_properties.title = "Instrukcja obsługi katalogu produktów w Nuxt Studio"
    doc.core_properties.subject = "Dodawanie, edycja i publikowanie produktów Termospec"
    doc.core_properties.author = "Termospec"
    doc.core_properties.keywords = "Nuxt Studio, Termospec, produkty, instrukcja"

    # Cover — customer_pack pattern, compact_reference_guide preset.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(34)
    if LOGO.exists():
        shape = p.add_run().add_picture(str(LOGO), width=Inches(1.65))
        add_alt_text(shape, "Logo Termospec", "Logo firmy Termospec")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run("INSTRUKCJA DLA REDAKTORA"), size=10, color=BRAND, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    set_run_font(p.add_run("Katalog produktów\nw Nuxt Studio"), name="Calibri", size=31, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(24)
    set_run_font(
        p.add_run("Dodawanie nowych produktów, edycja danych, media, status publikacji i bezpieczne wdrażanie zmian."),
        size=13.5,
        color=MUTED,
    )

    add_callout(
        doc,
        "Cel dokumentu",
        "umożliwić samodzielne uzupełnianie katalogu bez edytowania kodu strony i bez ryzyka przypadkowego opublikowania niezatwierdzonych danych.",
    )

    for label, value in (
        ("Dla", "redaktorów treści Termospec"),
        ("Zakres", "polska wersja katalogu produktów"),
        ("Wersja", "1.0 / 20 sierpnia 2026"),
        ("System", "Nuxt Studio 1.7 + logowanie firmowe + Vercel"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        set_run_font(p.add_run(f"{label}: "), size=10.5, bold=True, color=BRAND)
        set_run_font(p.add_run(value), size=10.5, color=INK)

    doc.add_heading("Spis instrukcji", level=2)
    add_bullets(doc, [
        "1. Dostęp i układ Nuxt Studio",
        "2. Edycja istniejącego produktu",
        "3. Znaczenie pól produktu",
        "4. Zdjęcia, schematy i dokumenty",
        "5. Dodawanie nowego produktu z JSON",
        "6. Kontrola i publikacja",
        "7. Typowe problemy i lista kontrolna",
    ])

    add_safe_page_break(doc)
    doc.add_heading("Najważniejsze w 60 sekund", level=1)
    add_numbered(doc, [
        "Wejdź na adres Studio przekazany przez administratora (adres strony zakończony /_studio) i zaloguj się firmowym adresem e-mail oraz otrzymanym hasłem.",
        "Jeżeli produkt już istnieje w katalogu, wybierz Treść → Products i edytuj istniejący rekord. Nie twórz jego duplikatu.",
        "Jeżeli dodajesz zupełnie nowy model lub typoszereg, użyj przycisku Import produktu w dolnej części podglądu strony.",
        "Uzupełnij treść i media, obserwując podgląd po prawej stronie. Nie wpisuj niepotwierdzonych parametrów technicznych.",
        "Pozostaw status draft do chwili akceptacji produktu. Dopiero gotowy produkt zmień na published.",
        "Kliknij Przeglądaj, sprawdź listę zmian, wpisz krótki opis i wybierz Opublikuj. Poczekaj na zakończenie wdrożenia.",
    ])

    doc.add_heading("Dwa różne rodzaje wersji roboczej", level=2)
    add_field_table(doc, [
        ("Draft w Studio", "Zmiana jest zapisana w przeglądarce, ale nie trafiła jeszcze do GitHub. Znika z procesu publikacji dopiero po użyciu opcji Opublikuj albo Cofnij zmiany."),
        ("Status produktu: draft", "Rekord może zostać opublikowany do GitHub, ale produkt nadal pozostaje ukryty dla zwykłych użytkowników strony. Jest widoczny tylko w podglądzie Studio z oznaczeniem szkicu."),
        ("Status produktu: published", "Po opublikowaniu zmian i zakończeniu wdrożenia produkt jest widoczny w publicznym katalogu."),
    ], header=("Pojęcie", "Znaczenie"))

    add_callout(doc, "Zasada bezpieczeństwa", "Najpierw przygotuj i sprawdź produkt jako draft. Status published ustawiaj dopiero po akceptacji danych, zdjęć i miejsca w katalogu.")

    add_page_break(doc)
    doc.add_heading("1. Dostęp i układ Nuxt Studio", level=1)
    doc.add_heading("Logowanie", level=2)
    add_numbered(doc, [
        "Otwórz adres strony zakończony /_studio. Dokładny adres produkcyjny powinien zostać przekazany razem z uprawnieniami.",
        "Wpisz wskazany firmowy adres e-mail oraz hasło przekazane przez administratora strony. Konto Google ani konto GitHub nie są potrzebne.",
        "Kliknij Zaloguj się. Po poprawnym zalogowaniu otworzy się katalog produktów z aktywnymi narzędziami Nuxt Studio.",
        "Nie zapisuj hasła na współdzielonym komputerze. Po zakończeniu pracy wyloguj się ze Studio i zamknij kartę przeglądarki.",
    ])
    add_callout(doc, "Brak dostępu", "Sprawdź pisownię całego firmowego adresu e-mail i wielkość liter w haśle. Po kilku błędnych próbach odczekaj 15 minut. Jeśli problem pozostaje, skontaktuj się z administratorem — nie zakładaj konta Google ani GitHub.")

    doc.add_heading("Układ ekranu", level=2)
    add_bullets(doc, [
        "Treść — kolekcje danych i pliki strony. Produkty znajdują się w kolekcji Products.",
        "Media — biblioteka obrazów i plików przechowywanych w katalogu publicznym.",
        "Podgląd strony — pokazuje efekt zmian na żywo, zanim zostaną opublikowane.",
        "Import produktu — dodatkowe narzędzie Termospec do tworzenia produktów z gotowego JSON.",
    ])
    add_safe_page_break(doc)
    add_screenshot(
        doc,
        "01-katalog-produktow.png",
        "Rysunek 1. Nuxt Studio otwarte nad katalogiem produktów: panel treści po lewej, nawigacja katalogu i podgląd produktu po prawej. Zrzut wykonano w środowisku szkoleniowym; w produkcji stopka panelu pokazuje połączenie z repozytorium.",
        "Widok Nuxt Studio z kolekcjami Product Translations, Products i Team oraz katalogiem produktów Termospec.",
        width=6.2,
    )

    add_page_break(doc)
    doc.add_heading("2. Edycja istniejącego produktu", level=1)
    add_callout(doc, "Rekomendowana ścieżka", "Dla pozycji już obecnych w katalogu zawsze edytuj istniejący rekord. Importer JSON służy przede wszystkim do nowych modeli albo do pierwszego, pełnego uzupełnienia rekordu przygotowanego poza Studio.")
    add_numbered(doc, [
        "W zakładce Treść wybierz kolekcję Products.",
        "Na liście znajdź produkt po nazwie lub technicznej nazwie pliku, np. agregaty-tlokowe-ts-150.yml.",
        "Kliknij kartę produktu. Po lewej otworzy się formularz, a po prawej pozostanie podgląd katalogu.",
        "Zmień potrzebne pola. Studio zapisuje je jako lokalne zmiany robocze — nie ma osobnego przycisku Zapisz.",
        "Sprawdź podgląd produktu oraz status. Opublikuj dopiero po zakończeniu całej korekty.",
    ])
    add_screenshot(
        doc,
        "02-lista-produktow.png",
        "Rysunek 2. Kolekcja Products zawiera osobny rekord dla każdego modelu lub typoszeregu.",
        "Lista rekordów produktów w Nuxt Studio z nazwami plików YAML.",
    )

    add_safe_page_break(doc)
    doc.add_heading("Jak czytać nazwę rekordu", level=2)
    add_body(doc, "Nazwa widoczna na karcie jest techniczną wersją nazwy pliku. Przykład: agregaty-tlokowe-ts-150.yml odpowiada produktowi TS-150. Rozszerzenie YAML lub JSON nie zmienia sposobu edycji — oba typy otwierają ten sam formularz.")
    add_bullets(doc, [
        "Nie zmieniaj nazwy pliku istniejącego produktu bez konsultacji z administratorem.",
        "Nie duplikuj rekordu tylko po to, aby zmienić tytuł produktu.",
        "Jeśli na liście są dwa podobne rekordy, przed edycją porównaj ich ścieżkę i pozycję katalogową.",
    ])

    add_safe_page_break(doc, extra_space=54)
    doc.add_heading("Formularz i podgląd produktu", level=2)
    add_body(doc, "Formularz produktu jest generowany automatycznie. Zmiany tekstu po lewej powinny pojawiać się w podglądzie po prawej. Jeżeli pole pozostanie puste, odpowiadająca mu sekcja nie będzie wyświetlana publicznie.")
    add_screenshot(
        doc,
        "03-edycja-produktu-ts150.png",
        "Rysunek 3. Edycja produktu TS-150: podstawowe pola, miejsce w katalogu, status oraz podgląd produktu.",
        "Formularz produktu TS-150 w Nuxt Studio z polami Title, Category, Catalogue Placement, Status, Lead i Short Description.",
        width=5.3,
    )
    add_callout(doc, "Ważne", "Nazwy niektórych pól są obecnie wyświetlane po angielsku przez Nuxt Studio. W kolejnych rozdziałach znajduje się polskie objaśnienie wszystkich pól używanych przy produktach.")

    add_page_break(doc)
    doc.add_heading("3. Znaczenie pól produktu", level=1)
    doc.add_heading("Tożsamość i pozycja w katalogu", level=2)
    add_field_table(doc, [
        ("Title", "Nazwa widoczna w katalogu, np. TS-150. Zmiana tytułu nie wymaga zmiany nazwy pliku."),
        ("Model Code", "Pełne oznaczenie techniczne, jeśli różni się od tytułu, np. TS-150/ZM. Pole opcjonalne."),
        ("Category", "Czytelna nazwa rodziny widoczna przy produkcie, np. Agregaty chłodnicze ze sprężarką tłokową."),
        ("Category Slug", "Główna kategoria katalogu. Dla istniejącego produktu nie zmieniaj jej bez sprawdzenia nawigacji."),
        ("Group Slug", "Podkategoria. Może być pusta, jeśli produkt należy bezpośrednio do kategorii głównej."),
        ("Kind", "model — pojedynczy produkt; series — typoszereg; description — opis możliwości, a nie konkretny model."),
        ("Status", "draft ukrywa produkt publicznie; published pokazuje go po publikacji i wdrożeniu."),
    ])

    doc.add_heading("Treść i SEO", level=2)
    add_field_table(doc, [
        ("Lead", "Najważniejsze jedno zdanie: czym jest produkt i do czego służy."),
        ("Short Description", "Drugie, krótkie streszczenie. Pozostaw puste, jeśli Lead wystarcza."),
        ("SEO / Title", "Tytuł dla wyszukiwarki. Jeśli pusty, system może użyć nazwy produktu."),
        ("SEO / Description", "Krótki i konkretny opis dla wyszukiwarki, bez listy wszystkich parametrów."),
        ("Power Range", "Moc pojedynczego modelu lub zakres całego typoszeregu, zawsze z jednostką."),
    ])

    add_callout(doc, "Nie wymyślaj danych", "Jeżeli wartość nie została potwierdzona w dokumentacji technicznej, pozostaw pole puste. Brak sekcji jest lepszy niż niezweryfikowany parametr.", fill="FFF1F3", accent=BRAND)

    add_safe_page_break(doc, extra_space=54)
    doc.add_heading("Parametry, listy i sekcje", level=2)
    add_field_table(doc, [
        ("Specs", "Maksymalnie kilka najważniejszych parametrów widocznych od razu. Wartość wpisuj razem z jednostką, np. 150 kW."),
        ("Technical Data", "Pełna tabela danych. Wartość i jednostka są osobnymi polami, np. 150 oraz kW."),
        ("Technical Notes", "Warunki pomiaru, tolerancje, wyjątki i zastrzeżenia do parametrów."),
        ("Applications", "Zastosowania, miejsca pracy i procesy, dla których przeznaczony jest produkt."),
        ("Components", "Najważniejsze podzespoły urządzenia lub elementy zestawu."),
        ("Features", "Zweryfikowane cechy techniczne i użytkowe."),
        ("Related Products", "Kompatybilne urządzenia lub produkty uzupełniające."),
        ("Service Scope", "Zakres doboru, uruchomienia, przeglądów, modernizacji i serwisu Termospec."),
        ("Sections", "Dłuższe bloki opisowe, np. Zastosowanie, Zasada działania, Konstrukcja. Każdy blok ma tytuł, treść i opcjonalne punkty."),
    ])

    add_page_break(doc)
    doc.add_heading("Obsługa list", level=2)
    add_bullets(doc, [
        "Dodaj items — dodaje nową pozycję do listy.",
        "Ikona edycji — otwiera treść istniejącej pozycji.",
        "Ikona usuwania — usuwa tylko wskazany element listy.",
        "Strzałki góra/dół — zmieniają kolejność parametrów, sekcji albo zdjęć.",
    ])
    add_callout(doc, "Dobra praktyka", "Jedna pozycja listy powinna zawierać jedną informację. Nie wklejaj kilku punktów w jednym polu i nie wpisuj ręcznie numeracji.")

    doc.add_heading("4. Zdjęcia, schematy i dokumenty", level=1)
    add_field_table(doc, [
        ("Hero Image", "Główny render albo zdjęcie konkretnego produktu. Nie używaj grafiki całej kategorii, jeśli rekord dotyczy jednego modelu."),
        ("Image Alt", "Krótki opis tego, co przedstawia główny obraz. Nie zaczynaj od słów „zdjęcie” lub „grafika”."),
        ("Header Sketch", "Opcjonalny schemat techniczny lub rysunek urządzenia, z osobnym opisem Alt."),
        ("Gallery", "Dodatkowe zdjęcia produktu, instalacji i detali. Ustaw kolejność od najważniejszego ujęcia."),
        ("Documents", "Karty katalogowe, instrukcje, certyfikaty i pliki PDF. Każdy plik powinien mieć czytelną nazwę."),
        ("Related Projects", "Powiązane realizacje z numerem, nazwą, opisem zakresu i materiałem ilustracyjnym."),
    ])
    doc.add_heading("Dodawanie plików", level=2)
    add_numbered(doc, [
        "Otwórz zakładkę Media w lewym górnym rogu Studio.",
        "Przejdź do właściwego folderu, najlepiej images/products dla materiałów produktowych.",
        "Prześlij plik albo przeciągnij go do biblioteki. Stosuj krótkie nazwy: małe litery, cyfry i myślniki, bez spacji.",
        "Wróć do produktu i wybierz plik w polu z selektorem mediów.",
        "Uzupełnij tekst alternatywny i sprawdź podgląd.",
    ])
    add_safe_page_break(doc)
    add_screenshot(
        doc,
        "06-biblioteka-mediow.png",
        "Rysunek 4. Zakładka Media pokazuje pliki z katalogu publicznego i umożliwia ich wybór w formularzu produktu.",
        "Biblioteka mediów Nuxt Studio z miniaturami obrazów Termospec.",
        width=5.5,
    )
    add_callout(doc, "Formaty", "Dla zdjęć używaj JPG lub WebP, dla grafik z przezroczystością PNG, a dla dokumentów PDF. Nie przesyłaj wielokrotnie tego samego pliku pod różnymi nazwami.")

    add_page_break(doc)
    doc.add_heading("5. Dodawanie nowego produktu z JSON", level=1)
    add_callout(doc, "Kiedy używać", "Użyj importera dla nowego modelu/typoszeregu albo wtedy, gdy pełny opis został przygotowany w ustandaryzowanym JSON. Dla istniejącego produktu wybierz zwykłą edycję w kolekcji Products.")
    add_numbered(doc, [
        "Otwórz Studio na stronie katalogu produktów.",
        "Kliknij Import produktu w dolnej części podglądu.",
        "Pozostań w zakładce Nowy produkt.",
        "Wybierz sposób przygotowania danych: skopiuj instrukcję dla AI, pobierz pusty szablon, pobierz JSON Schema albo wczytaj gotowy plik.",
        "Wklej do pola Dane produktu jeden poprawny obiekt JSON — bez komentarzy, bez bloku Markdown i bez dodatkowych pól.",
    ])
    add_screenshot(
        doc,
        "04-importer-json.png",
        "Rysunek 5. Importer nowego produktu: szablon, instrukcja dla AI, plik JSON i pole walidowane przed utworzeniem rekordu.",
        "Okno Import produktu z JSON z polem Dane produktu i przyciskami szablonu oraz walidacji.",
    )

    add_safe_page_break(doc, extra_space=54)
    doc.add_heading("Najprostszy workflow z AI", level=2)
    add_numbered(doc, [
        "Kliknij Kopiuj instrukcję dla AI.",
        "Wklej instrukcję do ChatGPT, a następnie dodaj zweryfikowane materiały źródłowe produktu.",
        "Poproś o zwrot wyłącznie JSON i sprawdź, czy AI nie dopisało niepotwierdzonych parametrów.",
        "Skopiuj odpowiedź do importera i użyj Tylko sprawdź.",
        "Jeśli walidacja zgłosi błąd, popraw wskazane pole zamiast usuwać strukturę szablonu.",
    ])

    doc.add_heading("Miejsce produktu w katalogu", level=2)
    add_body(doc, "Po wklejeniu JSON wybierz miejsce wyświetlania. Ustawienia z formularza importera mają pierwszeństwo przed wartościami wpisanymi w pliku.")
    add_field_table(doc, [
        ("Kategoria", "Jedna z głównych kategorii katalogu, np. Agregaty chłodnicze."),
        ("Podkategoria", "Rodzina produktu w ramach kategorii. Jeśli nie występuje, pozostaw Bez podkategorii."),
        ("Sposób wyświetlania", "Model, Typoszereg albo Opis. Najczęściej wybierany jest Model."),
        ("Pozycja katalogowa", "Dla pozycji zatwierdzonej już w katalogu wybierz dokładnie istniejący model. Nową pozycję wybieraj tylko dla rzeczywiście nowego produktu."),
        ("Nazwa pliku", "Unikalna nazwa techniczna, np. wentylator-osiowy-wo-400. Importer dopisze rozszerzenie .json."),
    ])
    add_page_break(doc)
    add_screenshot(
        doc,
        "05-ustawienia-katalogu.png",
        "Rysunek 6. Dolna część importera zawiera kategorię, podkategorię, sposób wyświetlania, pozycję katalogową i nazwę pliku.",
        "Pełny formularz importu produktu z ustawieniami miejsca w katalogu i przyciskiem Utwórz wersję roboczą.",
        width=5.9,
    )

    doc.add_heading("Walidacja i utworzenie produktu", level=2)
    add_numbered(doc, [
        "Kliknij Tylko sprawdź. Popraw wszystkie komunikaty, zanim utworzysz produkt.",
        "Kliknij Utwórz wersję roboczą. System zawsze tworzy nowy rekord ze statusem draft.",
        "Po utworzeniu Studio otworzy formularz nowego produktu. Uzupełnij media i pola, których nie było w materiałach źródłowych.",
        "Sprawdź produkt w podglądzie katalogu — jako draft powinien mieć oznaczenie szkicu i nie być widoczny poza Studio.",
    ])

    add_callout(doc, "Ochrona przed duplikatem", "Jeśli system zgłosi, że plik już istnieje, przerwij import i znajdź istniejący rekord w Products. Nie rozwiązuj problemu przez dopisywanie przypadkowego numeru do nazwy.", fill="FFF1F3", accent=BRAND)

    add_safe_page_break(doc, extra_space=54)
    doc.add_heading("6. Kontrola i publikacja", level=1)
    doc.add_heading("Kontrola przed ustawieniem published", level=2)
    add_bullets(doc, [
        "Nazwa i kod modelu są poprawne oraz zgodne z dokumentacją.",
        "Produkt znajduje się w prawidłowej kategorii i podkategorii.",
        "Najważniejsze parametry nie przeczą pełnej tabeli technicznej.",
        "Jednostki, zakresy, przecinki dziesiętne i znaki specjalne zostały sprawdzone.",
        "Główne zdjęcie przedstawia właściwy produkt, a tekst Alt opisuje obraz.",
        "Sekcje puste lub niepotwierdzone zostały usunięte albo pozostawione puste.",
        "Podgląd nie zawiera pustych nagłówków, błędnych obrazów ani powtórzeń.",
    ])

    heading = doc.add_heading("Publikowanie zmian", level=2)
    heading.paragraph_format.page_break_before = True
    add_numbered(doc, [
        "Zmień Status produktu z draft na published tylko wtedy, gdy produkt ma pojawić się publicznie.",
        "W pasku Studio wybierz Przeglądaj, aby otworzyć listę utworzonych, zaktualizowanych, przemianowanych i usuniętych plików.",
        "Sprawdź, czy lista obejmuje wyłącznie zamierzone zmiany. W razie pomyłki wróć do edycji albo użyj Cofnij zmiany dla konkretnego pliku.",
        "W polu wiadomości wpisz zwięzły opis, np. „Uzupełnienie danych i zdjęć produktu TS-200”.",
        "Kliknij Opublikuj. Studio zapisze zmiany do repozytorium GitHub, a Vercel rozpocznie nowe wdrożenie.",
        "Poczekaj na komunikat Wdrożenie zakończone, odśwież Studio i sprawdź publiczną stronę w osobnej karcie.",
    ])
    add_callout(doc, "Czas wdrożenia", "Zmiana nie musi być widoczna natychmiast po kliknięciu Opublikuj. Najpierw powstaje commit w GitHub, następnie Vercel buduje i wdraża nową wersję strony.")

    doc.add_heading("Czego nie publikować w jednym pakiecie", level=2)
    add_bullets(doc, [
        "przypadkowych zmian kilku niezwiązanych produktów,",
        "usunięć albo zmian nazw plików, których nie rozpoznajesz,",
        "danych technicznych bez źródła i akceptacji,",
        "nowej pozycji katalogowej utworzonej tylko dlatego, że nie znaleziono istniejącego rekordu.",
    ])

    add_safe_page_break(doc, extra_space=54)
    doc.add_heading("7. Typowe problemy", level=1)
    add_field_table(doc, [
        ("Nie widzę produktu publicznie", "Sprawdź Status = published, następnie czy zmiany zostały opublikowane i czy wdrożenie Vercel się zakończyło."),
        ("Produkt jest widoczny tylko w Studio", "Najczęściej jego Status nadal wynosi draft. To prawidłowe zachowanie wersji roboczej."),
        ("Importer zgłasza niepoprawny JSON", "Usuń komentarze, znaczniki ```json, tekst przed/po obiekcie i końcowe przecinki. Wklej dokładnie jeden obiekt JSON."),
        ("Importer zgłasza dodatkowe pola", "AI zmieniło strukturę. Wróć do pobranego szablonu i usuń wyłącznie pola spoza schematu."),
        ("Plik już istnieje", "Nie twórz kopii. Otwórz istniejący rekord w Products i edytuj go."),
        ("Produkt jest w złej kategorii", "Sprawdź Category Slug i Group Slug. Dla istniejącego produktu wykonaj zmianę ostrożnie i zweryfikuj nawigację."),
        ("Obraz się nie wyświetla", "Sprawdź, czy plik istnieje w Media, czy ścieżka zaczyna się od / oraz czy nazwa i rozszerzenie są identyczne."),
        ("Zmiany nie pojawiają się w podglądzie", "Odczekaj chwilę, zmień pole ponownie lub odśwież Studio. Przed odświeżeniem upewnij się, że zmiana widnieje w przeglądzie zmian."),
        ("Studio pokazuje konflikt", "Nie publikuj. Odśwież stronę po zakończeniu ostatniego wdrożenia; jeśli konflikt pozostaje, skontaktuj się z administratorem."),
    ], header=("Problem", "Co zrobić"))

    add_safe_page_break(doc, extra_space=54)
    doc.add_heading("Sytuacje wymagające kontaktu z administratorem", level=2)
    add_bullets(doc, [
        "brak dostępu do /_studio albo niedziałające firmowe dane logowania,",
        "konflikt repozytorium albo nieudana publikacja,",
        "potrzeba zmiany nazwy istniejącego pliku lub stałej pozycji katalogowej,",
        "konieczność usunięcia produktu, kategorii albo podkategorii,",
        "błędny build lub brak wdrożenia po poprawnym opublikowaniu zmian.",
    ])

    add_safe_page_break(doc, extra_space=108)
    doc.add_heading("Lista kontrolna redaktora", level=1)
    doc.add_heading("Przed utworzeniem", level=2)
    add_bullets(doc, [
        "Sprawdziłem, czy produkt nie istnieje już w kolekcji Products.",
        "Mam zweryfikowane materiały źródłowe i właściwe pliki graficzne.",
        "Znam kategorię, podkategorię i typ pozycji: model/typoszereg/opis.",
    ])
    doc.add_heading("Przed ustawieniem published", level=2)
    add_bullets(doc, [
        "Tekst nie zawiera danych wymyślonych przez AI.",
        "Parametry, wartości i jednostki są spójne.",
        "Puste sekcje nie pojawiają się w podglądzie.",
        "Zdjęcia są właściwe, a opisy Alt uzupełnione.",
        "Produkt jest we właściwym miejscu katalogu.",
    ])
    doc.add_heading("Przed kliknięciem Opublikuj", level=2)
    add_bullets(doc, [
        "Lista zmian obejmuje tylko zamierzone pliki.",
        "Nie ma przypadkowych usunięć ani zmian nazw.",
        "Wiadomość publikacji jasno opisuje zakres zmian.",
        "Po wdrożeniu sprawdzę produkt na stronie publicznej.",
    ])

    add_safe_page_break(doc, extra_space=54)
    add_callout(doc, "Wersja angielska", "Zakładka Wersja angielska w importerze tworzy oddzielny draft EN i nie modyfikuje polskiego produktu. Publiczna angielska wersja katalogu jest osobnym etapem — do czasu jej uruchomienia nie traktuj importu EN jako publikacji na stronie.")

    doc.add_heading("Źródła i aktualność", level=2)
    p = doc.add_paragraph()
    set_run_font(p.add_run("Instrukcja odzwierciedla stan projektu Termospec i Nuxt Studio 1.7 z dnia 20.08.2026. Oficjalne informacje o edycji formularzy: "))
    add_hyperlink(p, "nuxt.studio/content", "https://nuxt.studio/content")
    set_run_font(p.add_run("; o bibliotece mediów: "))
    add_hyperlink(p, "nuxt.studio/medias", "https://nuxt.studio/medias")
    set_run_font(p.add_run("; o publikowaniu do Git: "))
    add_hyperlink(p, "nuxt.studio/git-providers", "https://nuxt.studio/git-providers")
    set_run_font(p.add_run("."))

    add_callout(doc, "Ostatnia zasada", "Jeśli nie masz pewności, pozostaw produkt jako draft i poproś o weryfikację. Draft można bezpiecznie poprawić; błędny published trafia do publicznego katalogu po wdrożeniu.", fill="EEF3F8", accent=INK)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
